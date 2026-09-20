"""Reportes descargables en PDF del semáforo de posicionamiento."""
from __future__ import annotations

import base64
import io
import re
import unicodedata
from datetime import date
from pathlib import Path
from typing import Any, Optional
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

NAVY = colors.HexColor("#111b2b")
GRIS = colors.HexColor("#667085")
LINEA = colors.HexColor("#dce2ea")
TONOS = (
    ("negativo", "Notas negativas o críticas", "#b63a3a"),
    ("neutro", "Notas neutras", "#c49a2c"),
    ("positivo", "Notas positivas", "#2f855a"),
)
MAX_NOTAS = 60
MARGEN = 48

# (regular, negrita). Se usa la primera pareja que exista en el sistema para
# que acentos y símbolos se impriman bien; si no hay ninguna, Helvetica.
_FUENTES = (
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial Unicode.ttf", None),
    ("/Library/Fonts/Arial Unicode.ttf", None),
)


def _registrar_fuentes() -> tuple[str, str]:
    registradas = pdfmetrics.getRegisteredFontNames()
    if "InformeNormal" in registradas:
        return "InformeNormal", "InformeNegrita"
    for normal, negrita in _FUENTES:
        if not Path(normal).exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont("InformeNormal", normal))
            usar_negrita = negrita if negrita and Path(negrita).exists() else normal
            pdfmetrics.registerFont(TTFont("InformeNegrita", usar_negrita))
            pdfmetrics.registerFontFamily(
                "InformeNormal", normal="InformeNormal", bold="InformeNegrita",
                italic="InformeNormal", boldItalic="InformeNegrita",
            )
            return "InformeNormal", "InformeNegrita"
        except Exception:
            continue
    return "Helvetica", "Helvetica-Bold"


def _limpio(valor: Any, limite: Optional[int] = None) -> str:
    """Texto seguro para Paragraph: sin emojis, con espacios normalizados y escapado."""
    texto = re.sub(r"[\U00010000-\U0010ffff]", "", str(valor if valor is not None else ""))
    texto = re.sub(r"\s+", " ", texto).strip()
    if limite and len(texto) > limite:
        texto = texto[: limite - 1].rstrip() + "…"
    return escape(texto)


def _fecha(iso: Any) -> str:
    try:
        return date.fromisoformat(str(iso)[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return ""


def _estilos(normal: str, negrita: str) -> dict[str, ParagraphStyle]:
    tinta = colors.HexColor("#1f2937")
    return {
        "titulo": ParagraphStyle("titulo", fontName=negrita, fontSize=20, leading=25, textColor=NAVY, spaceAfter=4),
        "sub": ParagraphStyle("sub", fontName=normal, fontSize=10, leading=14, textColor=GRIS, spaceAfter=10),
        "h2": ParagraphStyle("h2", fontName=negrita, fontSize=12.5, leading=16, textColor=NAVY, spaceBefore=14, spaceAfter=6),
        "cuerpo": ParagraphStyle("cuerpo", fontName=normal, fontSize=10, leading=14.5, textColor=tinta, spaceAfter=6),
        "chico": ParagraphStyle("chico", fontName=normal, fontSize=8.5, leading=11.5, textColor=GRIS, spaceAfter=4),
        "nota": ParagraphStyle("nota", fontName=normal, fontSize=9.5, leading=13, textColor=tinta),
        "celda": ParagraphStyle("celda", fontName=normal, fontSize=9, leading=12, textColor=tinta),
        "celda_b": ParagraphStyle("celda_b", fontName=negrita, fontSize=9, leading=12, textColor=NAVY),
        "cifra": ParagraphStyle("cifra", fontName=negrita, fontSize=17, leading=20),
    }


def _imagen(data_url: Optional[str], ancho: float) -> Optional[Image]:
    """Decodifica la gráfica PNG enviada por el navegador; None si no es válida."""
    if not data_url or not data_url.startswith("data:image/png;base64,"):
        return None
    try:
        crudo = base64.b64decode(data_url.split(",", 1)[1], validate=True)
        if not crudo.startswith(b"\x89PNG"):
            return None
        w, h = ImageReader(io.BytesIO(crudo)).getSize()
        if not w or not h:
            return None
        alto = min(ancho * h / w, 270)
        return Image(io.BytesIO(crudo), width=alto * w / h, height=alto)
    except Exception:
        return None


def _tabla_semaforo(semaforo: dict, est: dict, ancho: float) -> Table:
    celdas = []
    for clave, etiqueta, color in TONOS:
        dato = semaforo.get(clave) or {}
        pct = float(dato.get("porcentaje") or 0)
        notas = int(dato.get("notas") or 0)
        celdas.append([
            Paragraph(f'<font color="{color}">{pct:.1f}%</font>', est["cifra"]),
            Paragraph(f"{etiqueta[6:].capitalize()} · {notas} {'nota' if notas == 1 else 'notas'}", est["celda"]),
        ])
    tabla = Table([celdas], colWidths=[ancho / 3] * 3)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f6f8fa")),
        ("BOX", (0, 0), (-1, -1), 0.5, LINEA),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ("LEFTPADDING", (0, 0), (-1, -1), 11),
    ]
    for i, (_, _, color) in enumerate(TONOS):
        estilo.append(("LINEABOVE", (i, 0), (i, 0), 3, colors.HexColor(color)))
    tabla.setStyle(TableStyle(estilo))
    return tabla


def _barra(semaforo: dict, ancho: float) -> Optional[Table]:
    partes = [(float((semaforo.get(k) or {}).get("porcentaje") or 0), c) for k, _, c in TONOS]
    partes = [(p, c) for p, c in partes if p > 0]
    if not partes:
        return None
    total = sum(p for p, _ in partes)
    tabla = Table([[""] * len(partes)], colWidths=[ancho * p / total for p, _ in partes], rowHeights=[8])
    tabla.setStyle(TableStyle([("BACKGROUND", (i, 0), (i, 0), colors.HexColor(c)) for i, (_, c) in enumerate(partes)]))
    return tabla


def _seccion(reporte: dict[str, Any], est: dict, ancho: float) -> list:
    nombre = _limpio(reporte["nombre"], 120)
    noticias = reporte.get("noticias") or []
    fuentes = len({str(n.get("fuente")) for n in noticias if n.get("fuente")})
    semaforo = reporte.get("semaforo") or {}
    analisis = reporte.get("analisis") or {}
    cobertura = reporte.get("cobertura_semanal") or {}

    historia: list = [
        Paragraph(nombre, est["titulo"]),
        Paragraph(
            f"Periodo del {_fecha(reporte.get('fecha_inicio'))} al {_fecha(reporte.get('fecha_fin'))} · "
            f"{len(noticias)} notas de {fuentes} fuentes", est["sub"]),
        Paragraph("Tono de la cobertura", est["h2"]),
        _tabla_semaforo(semaforo, est, ancho),
    ]
    barra = _barra(semaforo, ancho)
    if barra is not None:
        historia += [Spacer(1, 6), barra]
    if semaforo.get("aviso"):
        historia += [Spacer(1, 4), Paragraph(_limpio(semaforo["aviso"]), est["chico"])]

    historia += [Paragraph("Resumen de la cobertura", est["h2"]),
                 Paragraph(_limpio(analisis.get("resumen")) or "Sin resumen disponible.", est["cuerpo"])]
    por_tono = analisis.get("por_tono") or {}
    for clave, etiqueta, color in TONOS:
        texto = por_tono.get(clave)
        if texto:
            historia.append(Paragraph(f'<font color="{color}"><b>{etiqueta}.</b></font> {_limpio(texto)}', est["cuerpo"]))
    if analisis.get("lectura"):
        historia += [Paragraph("¿Qué representa esta cobertura?", est["h2"]),
                     Paragraph(_limpio(analisis["lectura"]), est["cuerpo"])]
    detalle = " · ".join(x for x in (
        f"Alcance: {_limpio(analisis.get('alcance'))}" if analisis.get("alcance") else "",
        f"Síntesis: {_limpio(analisis.get('metodo_resumen'))}" if analisis.get("metodo_resumen") else "",
    ) if x)
    if detalle:
        historia.append(Paragraph(detalle, est["chico"]))

    historia.append(Paragraph("Cobertura por semana", est["h2"]))
    imagen = _imagen(reporte.get("grafica_semanal"), ancho)
    if imagen is not None:
        historia.append(imagen)
    elif cobertura.get("con_fecha"):
        historia.append(Paragraph("La gráfica no estaba disponible al exportar el reporte.", est["chico"]))
    else:
        historia.append(Paragraph("Ninguna nota tiene una fecha de publicación identificable en el periodo.", est["chico"]))
    if cobertura.get("sin_fecha"):
        historia.append(Paragraph(
            f"{cobertura['sin_fecha']} notas no aparecen en la gráfica porque no se pudo identificar su fecha de publicación.",
            est["chico"]))

    asociaciones = (reporte.get("asociaciones") or [])[:15]
    if asociaciones:
        temas = ", ".join(f"{_limpio(a.get('frase'), 60)} ({int(a.get('noticias') or 0)})" for a in asociaciones)
        historia += [Paragraph("Palabras y frases más mencionadas", est["h2"]),
                     Paragraph(temas, est["cuerpo"]),
                     Paragraph("El número indica en cuántas notas distintas aparece cada expresión.", est["chico"])]

    historia.append(Paragraph("Notas consultadas", est["h2"]))
    colores = {clave: color for clave, _, color in TONOS}
    for nota in noticias[:MAX_NOTAS]:
        tono = str(nota.get("tono") or "neutro")
        color = colores.get(tono, "#c49a2c")
        titulo = _limpio(nota.get("titulo"), 220) or "Sin título"
        url = str(nota.get("url") or "")
        enlace = (f'<a href="{escape(url, {chr(34): "&quot;"})}" color="#175d80">{titulo}</a>'
                  if url.startswith(("http://", "https://")) else titulo)
        fecha = _fecha(nota.get("fecha")) or _limpio(nota.get("fecha_buscador")) or "Fecha no indicada"
        origen = "texto extraído" if nota.get("texto_disponible") else "solo título y resumen"
        historia.append(KeepTogether([
            Paragraph(f'<font color="{color}"><b>[{_limpio(tono)}]</b></font> {enlace}', est["nota"]),
            Paragraph(f"{_limpio(nota.get('fuente'))} · {fecha} · {origen}", est["chico"]),
            Spacer(1, 3),
        ]))
    if len(noticias) > MAX_NOTAS:
        historia.append(Paragraph(f"Se muestran {MAX_NOTAS} de {len(noticias)} notas.", est["chico"]))
    return historia


def _semana_pico(cobertura: dict) -> str:
    semanas = cobertura.get("semanas") or []
    if not semanas:
        return "—"
    mejor = max(semanas, key=lambda s: int(s.get("total") or 0))
    if not int(mejor.get("total") or 0):
        return "—"
    return f"Semana del {_fecha(mejor.get('semana'))} ({int(mejor['total'])} notas)"


def _comparativo(reportes: list[dict], est: dict, ancho: float) -> Table:
    def pct(r: dict, clave: str) -> str:
        return f"{float(((r.get('semaforo') or {}).get(clave) or {}).get('porcentaje') or 0):.1f} %"

    def temas(r: dict) -> str:
        return ", ".join(_limpio(a.get("frase"), 40) for a in (r.get("asociaciones") or [])[:3]) or "—"

    filas = [
        ("Notas", lambda r: str(len(r.get("noticias") or []))),
        ("Fuentes distintas", lambda r: str(len({str(n.get("fuente")) for n in r.get("noticias") or [] if n.get("fuente")}))),
        ("Negativas o críticas", lambda r: pct(r, "negativo")),
        ("Neutras", lambda r: pct(r, "neutro")),
        ("Positivas", lambda r: pct(r, "positivo")),
        ("Semana con más notas", lambda r: _semana_pico(r.get("cobertura_semanal") or {})),
        ("Expresiones principales", temas),
    ]
    datos = [[Paragraph("", est["celda"])] + [Paragraph(_limpio(r["nombre"], 80), est["celda_b"]) for r in reportes]]
    for etiqueta, funcion in filas:
        datos.append([Paragraph(etiqueta, est["celda_b"])] + [Paragraph(funcion(r), est["celda"]) for r in reportes])
    primera = 130
    resto = (ancho - primera) / len(reportes)
    tabla = Table(datos, colWidths=[primera] + [resto] * len(reportes), repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f3f7")),
        ("LINEBELOW", (0, 0), (-1, -1), 0.5, LINEA),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    return tabla


def nombre_archivo(reportes: list[dict[str, Any]]) -> str:
    def slug(texto: str) -> str:
        ascii_ = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
        return re.sub(r"[^A-Za-z0-9]+", "_", ascii_).strip("_").lower()[:40] or "persona"
    return "posicionamiento_" + "_vs_".join(slug(str(r["nombre"])) for r in reportes) + ".pdf"


def pdf_posicionamiento(reportes: list[dict[str, Any]]) -> bytes:
    """Genera el PDF de una persona o la comparación de dos (una sección por persona)."""
    if not reportes:
        raise ValueError("No hay reportes para exportar.")
    normal, negrita = _registrar_fuentes()
    est = _estilos(normal, negrita)
    ancho = letter[0] - 2 * MARGEN

    historia: list = []
    if len(reportes) > 1:
        historia += [
            Paragraph("Comparación de cobertura en medios", est["titulo"]),
            Paragraph(f"{_limpio(reportes[0]['nombre'], 80)} frente a {_limpio(reportes[1]['nombre'], 80)} · "
                      f"generado el {date.today():%d/%m/%Y}", est["sub"]),
            _comparativo(reportes, est, ancho),
            Spacer(1, 8),
            Paragraph("Las cifras describen la cobertura publicada en el periodo consultado. "
                      "Cada persona tiene su detalle en las páginas siguientes.", est["chico"]),
        ]
    for indice, reporte in enumerate(reportes):
        if indice:
            historia.append(PageBreak())
        elif len(reportes) > 1:
            historia.append(Spacer(1, 22))  # la primera persona continúa bajo la tabla comparativa
        historia += _seccion(reporte, est, ancho)

    def pie(lienzo, doc) -> None:
        lienzo.saveState()
        lienzo.setFont(normal, 7.5)
        lienzo.setFillColor(GRIS)
        lienzo.drawString(MARGEN, 30, "Describe la cobertura publicada; no mide opinión pública ni verifica las notas.")
        lienzo.drawRightString(letter[0] - MARGEN, 30, f"Página {doc.page}")
        lienzo.restoreState()

    buffer = io.BytesIO()
    titulo = "Posicionamiento en medios: " + " y ".join(str(r["nombre"]) for r in reportes)
    SimpleDocTemplate(
        buffer, pagesize=letter, leftMargin=MARGEN, rightMargin=MARGEN, topMargin=50, bottomMargin=54,
        title=titulo[:200], author="Inteligencia Electoral SLP",
    ).build(historia, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Discursos: Word, PDF y TXT
# ---------------------------------------------------------------------------

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _texto_plano(texto: str) -> str:
    """Quita restos de Markdown que el modelo pudiera haber añadido y caracteres no válidos en Word."""
    texto = _CONTROL.sub("", texto)
    texto = re.sub(r"\*\*(.+?)\*\*", r"\1", texto)
    texto = re.sub(r"(?m)^\s{0,3}#{1,6}\s*", "", texto)
    return texto.replace("`", "").strip()


def _parrafos(texto: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", _texto_plano(texto)) if p.strip()]


def notas_verificacion(verificacion: Optional[dict[str, Any]]) -> list[str]:
    """Texto para el pie de los documentos; se calcula en el servidor, nunca lo envía el navegador."""
    if not verificacion or verificacion.get("estado") in (None, "sin_cifras"):
        return []
    donde = verificacion.get("territorio")
    buenas, dudosas = verificacion.get("respaldadas", []), verificacion.get("no_respaldadas", [])
    notas = []
    if buenas:
        origen = f"los datos de {donde}" if donde else "lo indicado en la petición"
        notas.append(f"{len(buenas)} cifra(s) coinciden con {origen}.")
    if dudosas:
        listado = ", ".join(c["texto"] for c in dudosas[:12])
        notas.append(f"Pendientes de verificar ({len(dudosas)}): {listado}. No coinciden con los datos disponibles "
                     "ni con lo indicado en la petición; confírmalas antes de publicar.")
    return notas


def txt_discurso(texto: str) -> bytes:
    return _texto_plano(texto).encode("utf-8-sig")


def word_discurso(titulo: str, texto: str, verificacion: Optional[dict[str, Any]] = None) -> bytes:
    doc = Document()
    seccion = doc.sections[0]
    seccion.top_margin = seccion.bottom_margin = Inches(0.9)
    seccion.left_margin = seccion.right_margin = Inches(1.05)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11.5)

    doc.add_heading(_CONTROL.sub("", titulo) or "Discurso", 0)
    sello = doc.add_paragraph()
    run = sello.add_run(f"Borrador generado con asistencia de IA · {date.today():%d/%m/%Y}")
    run.italic, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(0x66, 0x70, 0x85)

    for parrafo in _parrafos(texto):
        bloque = doc.add_paragraph()
        bloque.paragraph_format.space_after = Pt(9)
        bloque.paragraph_format.line_spacing = 1.25
        lineas = parrafo.split("\n")
        for i, linea in enumerate(lineas):
            r = bloque.add_run(linea)
            if i < len(lineas) - 1:
                r.add_break()

    notas = notas_verificacion(verificacion)
    if notas:
        doc.add_heading("Verificación de cifras", level=2)
        for nota in notas:
            doc.add_paragraph(nota)
    aviso = doc.add_paragraph()
    r = aviso.add_run("Revisa el texto antes de usarlo: el borrador puede contener errores.")
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0x66, 0x70, 0x85)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def pdf_discurso(titulo: str, texto: str, verificacion: Optional[dict[str, Any]] = None) -> bytes:
    normal, negrita = _registrar_fuentes()
    est = _estilos(normal, negrita)
    est["lectura"] = ParagraphStyle("lectura", parent=est["cuerpo"], fontSize=11.5, leading=18, spaceAfter=12)
    historia: list = [
        Paragraph(_limpio(_texto_plano(titulo), 150) or "Discurso", est["titulo"]),
        Paragraph(f"Borrador generado con asistencia de IA · {date.today():%d/%m/%Y}", est["sub"]),
    ]
    for parrafo in _parrafos(texto):
        historia.append(Paragraph("<br/>".join(_limpio(l) for l in parrafo.split("\n")), est["lectura"]))
    notas = notas_verificacion(verificacion)
    if notas:
        historia.append(Paragraph("Verificación de cifras", est["h2"]))
        historia += [Paragraph(_limpio(n), est["cuerpo"]) for n in notas]
    historia.append(Paragraph("Revisa el texto antes de usarlo: el borrador puede contener errores.", est["chico"]))

    def pie(lienzo, doc) -> None:
        lienzo.saveState()
        lienzo.setFont(normal, 8)
        lienzo.setFillColor(GRIS)
        lienzo.drawRightString(letter[0] - MARGEN, 30, f"Página {doc.page}")
        lienzo.restoreState()

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=letter, leftMargin=58, rightMargin=58, topMargin=56, bottomMargin=56,
                      title=titulo[:200], author="Inteligencia Electoral SLP").build(historia, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()
