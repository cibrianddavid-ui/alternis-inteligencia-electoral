"""Identifica distritos disputados y prepara documentos descargables."""
from __future__ import annotations

import io
import re
from collections import Counter
from typing import Any

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from xml.sax.saxutils import escape

EXCLUIDOS = {'LISTA_NOMINAL', 'TOTAL_VOTOS_CALCULADOS', 'NULOS', 'NO_REGISTRADAS'}


def datos_distrito(df: pd.DataFrame, mapa: dict[str, str], municipio: str, año: str = '') -> dict[str, Any]:
    requeridas = ('municipio', 'id_distrito_local', 'tipo_eleccion', 'partido', 'votos')
    faltantes = [x for x in requeridas if x not in mapa]
    if faltantes:
        raise ValueError('Faltan columnas para identificar el distrito: ' + ', '.join(faltantes))
    municipio_col = df[mapa['municipio']].astype(str).str.strip()
    subset = df.loc[municipio_col.str.casefold() == municipio.strip().casefold()].copy()
    if subset.empty:
        raise ValueError(f'No hay resultados para el municipio {municipio}.')
    eleccion = subset[mapa['tipo_eleccion']].astype(str).str.upper()
    subset = subset.loc[eleccion.str.contains(r'DIPUT|DIP_', regex=True) & eleccion.str.contains('LOC')]
    if subset.empty:
        raise ValueError('No hay resultados de diputación local para ese municipio.')
    if 'anio' in mapa:
        anios = subset[mapa['anio']].astype(str).str.strip()
        año = año or max(anios.unique(), key=lambda v: int(v) if v.isdigit() else v)
        subset = subset.loc[anios == año]
    partido = subset[mapa['partido']].astype(str).str.strip().str.upper()
    subset = subset.loc[~partido.isin(EXCLUIDOS)]
    subset = subset.loc[subset[mapa['id_distrito_local']].notna()]
    tabla = (subset.groupby([mapa['id_distrito_local'], mapa['partido']], dropna=False)[mapa['votos']]
             .sum().reset_index())
    opciones = []
    for distrito, bloque in tabla.groupby(mapa['id_distrito_local']):
        puestos = bloque.sort_values(mapa['votos'], ascending=False).head(2)
        if len(puestos) < 2:
            continue
        primero, segundo = puestos.iloc[0], puestos.iloc[1]
        margen = int(primero[mapa['votos']] - segundo[mapa['votos']])
        total = float(bloque[mapa['votos']].sum())
        opciones.append({'distrito': str(distrito), 'ganador': str(primero[mapa['partido']]),
                         'segundo': str(segundo[mapa['partido']]), 'votos_ganador': int(primero[mapa['votos']]),
                         'votos_segundo': int(segundo[mapa['votos']]), 'margen_votos': margen,
                         'margen_porcentaje': round(100 * margen / total, 2) if total else 0,
                         'votos_opciones': int(total)})
    if not opciones:
        raise ValueError('No hay al menos dos opciones con votos para comparar distritos.')
    seleccionado = min(opciones, key=lambda x: (x['margen_porcentaje'], x['margen_votos']))
    return {'municipio': municipio, 'anio': año or 'sin columna de año',
            'alcance': 'Votos registrados en el municipio para cada distrito local',
            'criterio': 'Menor diferencia porcentual entre primer y segundo lugar, respecto a los votos de las opciones registradas',
            'distritos_comparados': len(opciones), **seleccionado}


def redactar(peticion: str, evidencia: dict, noticias: list[dict]) -> str:
    """Ficha descriptiva basada en los datos disponibles, sin persuasión."""
    d = evidencia
    párrafos = [
        f"Para la diputación local en el municipio de {d['municipio']}, "
        f"el distrito local {d['distrito']} tuvo la menor diferencia porcentual "
        f"entre el primero y el segundo lugar en los datos disponibles de {d['anio']}.",
        f"En la parte del distrito registrada para el municipio, {d['ganador']} "
        f"sumó {d['votos_ganador']:,} votos y {d['segundo']} sumó "
        f"{d['votos_segundo']:,}. La diferencia fue de {d['margen_votos']:,} "
        f"votos, equivalente a {d['margen_porcentaje']} % de los votos registrados "
        f"para las opciones políticas en esa zona.",
        f"Se compararon {d['distritos_comparados']} distritos con al menos dos "
        "opciones políticas registradas. La comparación se limita a los registros "
        "del municipio; si un distrito abarca otros municipios, estos datos no "
        "describen su resultado completo.",
    ]
    if noticias:
        párrafos.append("Notas disponibles en el análisis de posicionamiento: " +
                         "; ".join(f"{n['titulo']} ({n['fuente']})" for n in noticias) +
                         ". Sus títulos se presentan como referencias, sin verificar sus afirmaciones.")
    return "\n\n".join(párrafos)


def word(titulo: str, cuerpo: str, evidencia: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.8)
    section.bottom_margin = Inches(.8)
    doc.add_heading(titulo, 0)
    for parrafo in cuerpo.split('\n'):
        if parrafo.strip():
            doc.add_paragraph(parrafo.strip())
    doc.add_heading('Datos utilizados', level=2)
    doc.add_paragraph(f"Distrito local {evidencia['distrito']} · Municipio de {evidencia['municipio']} · Elección {evidencia['anio']}")
    doc.add_paragraph(f"Primer lugar: {evidencia['ganador']} ({evidencia['votos_ganador']:,} votos). Segundo: {evidencia['segundo']} ({evidencia['votos_segundo']:,} votos). Margen: {evidencia['margen_votos']:,} votos ({evidencia['margen_porcentaje']}%).")
    doc.add_paragraph(evidencia['alcance'] + '. ' + evidencia['criterio'] + '.')
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def pdf(titulo: str, cuerpo: str, evidencia: dict) -> bytes:
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Discurso', parent=styles['BodyText'], fontName='Helvetica',
                              fontSize=11, leading=16, spaceAfter=12))
    story = [Paragraph(escape(titulo), styles['Title']), Spacer(1, 18)]
    story += [Paragraph(escape(p.strip()), styles['Discurso']) for p in cuerpo.split('\n') if p.strip()]
    story += [Spacer(1, 14), Paragraph('Datos utilizados', styles['Heading2'])]
    story += [Paragraph(escape(f"Distrito local {evidencia['distrito']} · {evidencia['municipio']} · {evidencia['anio']}"), styles['Discurso'])]
    story += [Paragraph(escape(f"{evidencia['ganador']}: {evidencia['votos_ganador']:,} votos; {evidencia['segundo']}: {evidencia['votos_segundo']:,} votos. Margen: {evidencia['margen_votos']:,} votos ({evidencia['margen_porcentaje']}%)."), styles['Discurso'])]
    story += [Paragraph(escape(evidencia['alcance'] + '. ' + evidencia['criterio'] + '.'), styles['Discurso'])]
    SimpleDocTemplate(buffer, pagesize=letter, leftMargin=55, rightMargin=55).build(story)
    return buffer.getvalue()
